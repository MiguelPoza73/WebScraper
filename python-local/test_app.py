import csv
import io
import unittest
from unittest.mock import patch
import requests
import pandas as pd

from bs4 import BeautifulSoup
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest
from scraper import ScrapingError, extract_table, extract_css, fetch_page, find_tables, validate_url, _request
from exporter import to_csv, to_excel

HTML = '<table><tr><th> Nombre </th><th>Nombre</th></tr><tr><td> 001 </td><td></td></tr><tr><td> 001 </td><td>=1+1</td></tr></table>'


class ExtractionTests(unittest.TestCase):
    def setUp(self):
        self.soup = BeautifulSoup(HTML, "html.parser")
        self.data = extract_table(find_tables(self.soup)[0])

    def test_preserve_text_and_duplicates(self):
        self.assertEqual(list(self.data.columns), [" Nombre ", "Nombre"])
        self.assertEqual(self.data.values.tolist(), [[" 001 ", ""], [" 001 ", "=1+1"]])

    def test_csv_roundtrip(self):
        rows = list(csv.reader(io.StringIO(to_csv(self.data).decode("utf-8-sig"))))
        self.assertEqual(rows, [[" Nombre ", "Nombre"], [" 001 ", ""], [" 001 ", "=1+1"]])

    def test_excel_literal_formula(self):
        sheet = load_workbook(io.BytesIO(to_excel(self.data))).active
        self.assertEqual(sheet.cell(2, 1).value, " 001 ")
        self.assertEqual(sheet.cell(3, 2).value, "=1+1")
        self.assertEqual(sheet.cell(3, 2).data_type, "s")
        self.assertEqual(sheet.max_row, 3)

    def test_css_missing_and_links(self):
        soup = BeautifulSoup('<p><a href="/x"> A </a></p><p></p>', "html.parser")
        data = extract_css(soup, "p", [("text", "a", ""), ("link", "a", "href")])
        self.assertEqual(data.values.tolist(), [[" A ", "/x"], [None, None]])
        self.assertTrue(extract_css(soup, "h2", [("t", "", "")]).empty)

    def test_css_invalid_or_ambiguous(self):
        for selector in ("[",):
            with self.assertRaises(ScrapingError):
                extract_css(self.soup, selector, [("x", "", "")])
        with self.assertRaises(ScrapingError):
            extract_css(self.soup, "table", [("x", "td", "")])

    def test_complex_table_rejected(self):
        for html in ('<table><tr><td colspan="2">a</td></tr></table>', '<table><tr><td>a</td></tr><tr><td>b</td><td>c</td></tr></table>'):
            with self.assertRaises(ScrapingError):
                extract_table(BeautifulSoup(html, "html.parser").table)

    def test_multi_headers(self):
        data = extract_table(BeautifulSoup('<table><tr><th>A</th></tr><tr><th>B</th></tr><tr><td>001</td></tr></table>', 'html.parser').table)
        sheet = load_workbook(io.BytesIO(to_excel(data))).active
        self.assertEqual([sheet.cell(i, 1).value for i in range(1, 4)], ["A", "B", "001"])

    def test_bad_url(self):
        for url in ("file:///secret", "https://user:pass@example.com", "http://127.0.0.1", "http://["):
            with self.assertRaises(ScrapingError):
                validate_url(url)

    def test_excel_rejects_truncation(self):
        with self.assertRaisesRegex(ValueError, "32.767"):
            to_excel(pd.DataFrame([["a" * 32768]]))

    @patch("scraper.validate_url", side_effect=lambda url: url)
    @patch("scraper.time.sleep")
    @patch("scraper.requests.get", side_effect=requests.Timeout)
    def test_connection_timeout(self, *_):
        with self.assertRaisesRegex(ScrapingError, "tiempo"):
            _request("https://example.com")

    @patch("scraper.validate_url", side_effect=lambda url: url)
    @patch("scraper._request")
    def test_robots_disallow(self, request, _):
        request.return_value = (200, {}, b"User-agent: *\nDisallow: /\n")
        with self.assertRaises(ScrapingError):
            fetch_page("https://example.com/a")
        self.assertEqual(request.call_count, 1)

    @patch("scraper.validate_url", side_effect=lambda url: url)
    @patch("scraper._request")
    def test_http_failure(self, request, _):
        request.side_effect = [(404, {}, b""), (429, {}, b"")]
        with self.assertRaisesRegex(ScrapingError, "429"):
            fetch_page("https://example.com/a")

    def test_interface_load_and_extract(self):
        with patch("scraper.fetch_page", return_value=(self.soup, "https://example.com")):
            app = AppTest.from_file("main.py", default_timeout=30).run()
            app.checkbox[0].check()
            app.button[0].click().run()
            app.button[1].click().run()
            self.assertFalse(app.exception)
            self.assertEqual(app.session_state["result"].values.tolist(), self.data.values.tolist())


if __name__ == "__main__":
    unittest.main()
