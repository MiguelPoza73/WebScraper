"""Comprobación real: sin índices de tabla fijos ni resultados simulados."""
from io import BytesIO
from openpyxl import load_workbook
from scraper import EXAMPLE, ScrapingError, fetch_page, find_tables, extract_table
from exporter import to_csv, to_excel

if __name__ == "__main__":
    try:
        soup, url = fetch_page(EXAMPLE)
        tables = find_tables(soup)
        print(f"Fuente: {url}\nTablas encontradas: {len(tables)}")
        candidates = [table for table in tables if "Population" in table.get_text() and "India" in table.get_text()]
        if not candidates:
            raise ScrapingError("No se encontró la tabla de población: comprueba la estructura actual.")
        data = extract_table(candidates[0])
        sheet = load_workbook(BytesIO(to_excel(data))).active
        assert sheet.max_column == len(data.columns)
        assert to_csv(data)
        print(f"Extracción y exportación verificadas: {len(data)} filas, {len(data.columns)} columnas.")
    except ScrapingError as exc:
        print(f"Comprobación real no completada: {exc}")
        raise SystemExit(1)
